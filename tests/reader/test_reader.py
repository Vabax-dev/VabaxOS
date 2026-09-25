#!/usr/bin/env python3
"""Tests for vabaxos-reader: documents, reading position, and the window.

The window runs on a hidden Broadway display; with the Kokoro model
(VABAXOS_KOKORO_DIR) it also reads the first sentence, into a file
(VABAXOS_VOICE_RECORD), so nothing is heard.

    python3 tests/reader/test_reader.py
"""

import os
import shutil
import socket
import subprocess
import sys
import tempfile
import time
import unittest

REPO = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
ROOT = os.path.join(REPO, "packages", "vabaxos-reader", "root")
LIB = os.path.join(ROOT, "usr", "lib", "python3", "dist-packages")
VOICE_LIB = os.path.join(REPO, "packages", "vabaxos-voice", "root", "usr", "lib", "python3", "dist-packages")
PROGRAM = os.path.join(ROOT, "usr", "bin", "vabaxos-reader")
DATA = os.environ.get("VABAXOS_KOKORO_DIR", "")
sys.path.insert(0, LIB)

from vabaxos_reader.document import Document, ReadingState, read_paragraphs  # noqa: E402

TEXT = "Primo paragrafo. Seconda frase del primo paragrafo.\n\nSecondo paragrafo, una sola frase.\n"


def have(module):
    try:
        __import__(module)
        return True
    except ImportError:
        return False


class DocumentTest(unittest.TestCase):

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()

    def tearDown(self):
        self.tmp.cleanup()

    def path(self, name):
        return os.path.join(self.tmp.name, name)

    def test_text(self):
        with open(self.path("a.txt"), "w", encoding="utf-8") as f:
            f.write(TEXT)
        doc = Document.open(self.path("a.txt"))
        self.assertEqual(len(doc.paragraphs), 2)
        self.assertEqual([s[3] for s in doc.sentences],
                         ["Primo paragrafo.", "Seconda frase del primo paragrafo.", "Secondo paragrafo, una sola frase."])
        for _p, start, end, text in doc.sentences:
            self.assertEqual(doc.text[start:end], text)
        self.assertEqual(doc.next_paragraph(0, 1), 2)
        self.assertEqual(doc.sentence_at(doc.sentences[1][1] + 3), 1)

    def test_html(self):
        with open(self.path("a.html"), "w", encoding="utf-8") as f:
            f.write("<html><head><title>x</title><style>p{}</style></head><body><h1>Titolo</h1>"
                    "<p>Uno <b>due</b> tre.</p><script>no()</script><li>Voce</li></body></html>")
        self.assertEqual(read_paragraphs(self.path("a.html")), ["Titolo", "Uno due tre.", "Voce"])

    @unittest.skipUnless(have("docx"), "python3-docx is not installed")
    def test_docx(self):
        import docx
        d = docx.Document()
        d.add_heading("Titolo", 1)
        d.add_paragraph("Testo del documento Word.")
        d.save(self.path("a.docx"))
        self.assertEqual(read_paragraphs(self.path("a.docx")), ["Titolo", "Testo del documento Word."])

    @unittest.skipUnless(shutil.which("pdftotext") and have("PIL"), "pdftotext or python3-pil missing")
    def test_pdf_without_text_is_explained(self):
        from PIL import Image
        Image.new("RGB", (100, 100), "white").save(self.path("scan.pdf"))
        with self.assertRaisesRegex(ValueError, "scanned"):
            read_paragraphs(self.path("scan.pdf"))

    @unittest.skipUnless(have("ebooklib"), "python3-ebooklib is not installed")
    def test_epub(self):
        from ebooklib import epub
        book = epub.EpubBook()
        book.set_identifier("x")
        book.set_title("Libro")
        book.set_language("it")
        chapter = epub.EpubHtml(title="Uno", file_name="uno.xhtml", lang="it")
        chapter.content = "<h1>Capitolo uno</h1><p>C'era una volta.</p>"
        book.add_item(chapter)
        book.spine = [chapter]
        book.add_item(epub.EpubNcx())
        book.add_item(epub.EpubNav())
        epub.write_epub(self.path("a.epub"), book)
        self.assertEqual(read_paragraphs(self.path("a.epub")), ["Capitolo uno", "C'era una volta."])

    def scan(self, name):
        """A picture of an Italian sentence, as a photo or a scanned PDF."""
        from PIL import Image, ImageDraw, ImageFont
        image = Image.new("RGB", (1400, 300), "white")
        draw = ImageDraw.Draw(image)
        draw.text((40, 100), "Buongiorno, questa lettera arriva dal Comune.", fill="black",
                  font=ImageFont.load_default(size=56))
        image.save(self.path(name))
        return self.path(name)

    @unittest.skipUnless(shutil.which("tesseract") and have("PIL"), "tesseract or python3-pil missing")
    def test_image_is_recognized(self):
        from vabaxos_reader.document import NeedsOCR
        path = self.scan("lettera.png")
        with self.assertRaises(NeedsOCR):
            read_paragraphs(path)
        text = " ".join(read_paragraphs(path, ocr=True))
        self.assertIn("lettera arriva dal Comune", text)

    @unittest.skipUnless(shutil.which("tesseract") and shutil.which("pdftoppm") and have("PIL"),
                         "tesseract, pdftoppm or python3-pil missing")
    def test_scanned_pdf_is_recognized(self):
        path = self.scan("lettera.pdf")
        text = " ".join(read_paragraphs(path, ocr=True))
        self.assertIn("Buongiorno", text)

    def test_position_and_bookmarks(self):
        state = ReadingState(self.path("state.json"))
        state.set_position("/doc.txt", 5)
        state.add_bookmark("/doc.txt", 9, "nove")
        state.add_bookmark("/doc.txt", 3, "tre")
        again = ReadingState(self.path("state.json"))
        self.assertEqual(again.position("/doc.txt"), 5)
        self.assertEqual([m["index"] for m in again.bookmarks("/doc.txt")], [3, 9])

    def test_unsupported(self):
        with open(self.path("a.xyz"), "w", encoding="utf-8") as f:
            f.write("x")
        with self.assertRaises(ValueError):
            read_paragraphs(self.path("a.xyz"))


CHILD = r'''
import importlib.machinery, importlib.util, sys, time
loader = importlib.machinery.SourceFileLoader("reader", sys.argv[1])
spec = importlib.util.spec_from_loader("reader", loader)
reader = importlib.util.module_from_spec(spec)
loader.exec_module(reader)
from gi.repository import Adw, GLib, Gtk

def walk(widget):
    yield widget
    child = widget.get_first_child()
    while child is not None:
        yield from walk(child)
        child = child.get_next_sibling()

app = reader.ReaderApp()
def check(app):
    window = reader.ReaderWindow(app)
    window.present()
    # Buttons inside GTK's own widgets (window controls, drop-down, spin
    # button) get their accessible names from GTK; hidden buttons do not count.
    inner = (Gtk.WindowControls, Gtk.DropDown, Gtk.SpinButton)
    unnamed = [type(w).__name__ for w in walk(window) if isinstance(w, Gtk.Button) and w.get_mapped()
               and not w.get_label() and not w.get_tooltip_text()
               and not any(w.get_ancestor(kind) for kind in inner)]
    print("UNNAMED:" + ",".join(unnamed), flush=True)
    window.open_path(sys.argv[2])
    print("SENTENCES:%d" % len(window.document.sentences), flush=True)
    def wait_voice():
        if window.reader.kokoro is None and sys.argv[3] == "voice":
            return True
        if sys.argv[3] == "voice":
            window.toggle()
            GLib.timeout_add(8000, finish)
        else:
            finish()
        return False
    def finish():
        print("READING:%s" % window.reading, flush=True)
        window.stop_reading()
        app.quit()
        return False
    GLib.timeout_add(500, wait_voice)
app.connect("activate", check)
app.run([])
'''


def free_display():
    for number in range(60, 100):
        with socket.socket() as s:
            if s.connect_ex(("127.0.0.1", 8080 + number)) != 0:
                return number
    raise RuntimeError("no free Broadway display")


@unittest.skipUnless(shutil.which("gtk4-broadwayd"), "gtk4-broadwayd (libgtk-4-bin) is not installed")
class WindowTest(unittest.TestCase):

    def test_window_opens_names_and_reads(self):
        display = free_display()
        broadway = subprocess.Popen(["gtk4-broadwayd", f":{display}"], stdout=subprocess.DEVNULL,
                                    stderr=subprocess.DEVNULL)
        time.sleep(1)
        with tempfile.TemporaryDirectory() as tmp:
            doc = os.path.join(tmp, "a.txt")
            with open(doc, "w", encoding="utf-8") as f:
                f.write(TEXT)
            record = os.path.join(tmp, "audio.raw")
            voice = "voice" if DATA and os.path.exists(os.path.join(DATA, "kokoro-v1.0.onnx")) else "novoice"
            env = dict(os.environ, GDK_BACKEND="broadway", BROADWAY_DISPLAY=f":{display}",
                       GSETTINGS_BACKEND="memory", XDG_DATA_HOME=tmp, NO_AT_BRIDGE="1",
                       PYTHONPATH=f"{LIB}:{VOICE_LIB}", VABAXOS_VOICE_RECORD=record,
                       LANGUAGE="C", LANG="C.UTF-8", PYTHONDONTWRITEBYTECODE="1")
            result = subprocess.run([sys.executable, "-c", CHILD, PROGRAM, doc, voice], env=env,
                                    capture_output=True, text=True, timeout=120)
            broadway.terminate()
            broadway.wait()
            out = dict(line.split(":", 1) for line in result.stdout.splitlines() if ":" in line)
            self.assertEqual(out.get("UNNAMED"), "", result.stderr[-2000:])
            self.assertEqual(out.get("SENTENCES"), "3", result.stderr[-2000:])
            if voice == "voice":
                self.assertGreater(os.path.getsize(record), 24000, "no speech recorded")
                # Position saved at pause.
                self.assertTrue(os.path.exists(os.path.join(tmp, "vabaxos-reader", "state.json")))


if __name__ == "__main__":
    unittest.main(verbosity=2)
